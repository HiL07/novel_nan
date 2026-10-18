# -*- coding: utf-8 -*-
"""Reframe a harem / multi-love-interest web novel so it passes 番茄小说 review.

Default behaviour applies the bundled references/example_map.json (the proven
map for 《重生2026：开局四个绝色老婆，我靠AI赢麻了》, refactored to a
single-spouse + partners structure). Override per novel with --map / --title-*.

What it does (in order):
  1. Bulk substring replacements (title, explicit violations, neutralised
     body language, de-romanticised subtext) — preserves paragraph formatting
     by collapsing to the first run.
  2. Full-chapter swaps (e.g. a polygamy wedding chapter -> compliant version).
  3. Sets the document core title property.

Usage:
    python reframe_novel.py <input.docx> <output.docx> [--map map.json]
                            [--title-old OLD] [--title-new NEW]

map.json schema:
    {
      "title_old": "...", "title_new": "...",
      "bulk": [ ["old substring", "new substring"], ... ],
      "chapter_swaps": { "Chapter Title": {"next": "Next Chapter Title",
                                           "body": ["line", ...]} }
    }
"""
from errkit import safe_main
import argparse
import copy
import json
import os
import sys

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

HERE = os.path.dirname(os.path.abspath(__file__))
DEFAULT_MAP = os.path.join(HERE, "..", "references", "example_map.json")


def load_map(path):
    with open(path, encoding="utf-8") as f:
        return json.load(f)


def apply_bulk(text, bulk):
    for old, new in bulk:
        if old and old in text:
            text = text.replace(old, new)
    return text


def make_para(text, style_p):
    p = OxmlElement("w:p")
    if style_p._p.pPr is not None:
        p.append(copy.deepcopy(style_p._p.pPr))
    r = OxmlElement("w:r")
    if style_p.runs and style_p.runs[0]._r.rPr is not None:
        r.append(copy.deepcopy(style_p.runs[0]._r.rPr))
    t = OxmlElement("w:t")
    t.set(qn("xml:space"), "preserve")
    t.text = text
    r.append(t)
    p.append(r)
    return p


def apply_chapter_swaps(doc, swaps):
    for title_text, spec in swaps.items():
        paras = doc.paragraphs
        try:
            title_idx = next(i for i, p in enumerate(paras)
                             if p.text.strip() == title_text)
            next_idx = next(i for i, p in enumerate(paras)
                            if p.text.strip().startswith(spec["next"]))
        except StopIteration:
            print(f"  [warn] chapter swap skipped, title not found: {title_text}")
            continue
        for i in range(next_idx - 1, title_idx, -1):
            paras[i]._p.getparent().remove(paras[i]._p)
        prev = doc.paragraphs[title_idx]._p
        for line in spec["body"]:
            np = make_para(line, doc.paragraphs[title_idx])
            prev.addnext(np)
            prev = np
        print(f"  [ok] chapter swapped: {title_text}")


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("input")
    ap.add_argument("output")
    ap.add_argument("--map", default=DEFAULT_MAP)
    ap.add_argument("--title-old", default=None)
    ap.add_argument("--title-new", default=None)
    args = ap.parse_args()

    mp = load_map(args.map)
    title_old = args.title_old or mp.get("title_old", "")
    title_new = args.title_new or mp.get("title_new", "")
    bulk = mp.get("bulk", [])
    swaps = mp.get("chapter_swaps", {})

    doc = Document(args.input)
    n_changed = 0

    def rewrite(paras):
        nonlocal n_changed
        for p in paras:
            t = p.text
            new = apply_bulk(t, bulk)
            if new != t:
                if p.runs:
                    p.runs[0].text = new
                    for r in p.runs[1:]:
                        r.text = ""
                else:
                    p.add_run(new)
                n_changed += 1

    rewrite(doc.paragraphs)
    for tbl in doc.tables:
        for row in tbl.rows:
            for cell in row.cells:
                rewrite(cell.paragraphs)

    if swaps:
        apply_chapter_swaps(doc, swaps)

    if title_old and title_new:
        # also fix any in-body occurrences of the title string
        for p in doc.paragraphs:
            if title_old in p.text:
                p.runs[0].text = p.text.replace(title_old, title_new)
                for r in p.runs[1:]:
                    r.text = ""
        try:
            doc.core_properties.title = title_new
        except Exception:
            pass

    doc.save(args.output)
    print(f"reframe done. paragraphs changed: {n_changed} -> {args.output}")


if __name__ == "__main__":
    safe_main(main)

